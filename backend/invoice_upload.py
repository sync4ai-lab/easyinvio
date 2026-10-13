"""Bounded statement parsing and reproducible comparisons against server snapshots."""
import asyncio
import csv
import hashlib
import io
import json
import re
import zipfile
from decimal import Decimal
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError
from pathlib import PurePath
from typing import Any
from uuid import UUID

import asyncpg
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from pydantic import BaseModel, ConfigDict, Field, field_validator

from backend.ai_engine import strict_decimal

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
HEADERS = {'upstream_shipment_id', 'currency', 'invoice_weight_kg', 'wholesale_cost', 'residential_fee'}
HEADER_ALIASES = {
    'shipment_id': 'upstream_shipment_id', 'shipment_reference': 'upstream_shipment_id',
    'carrier_reference': 'upstream_shipment_id', 'carrier_reference_code': 'upstream_shipment_id',
    'invoice_currency': 'currency', 'currency_code': 'currency',
    'billed_weight_kg': 'invoice_weight_kg', 'invoice_weight_in_kg': 'invoice_weight_kg', 'weight_kg': 'invoice_weight_kg',
    'invoice_wholesale_cost': 'wholesale_cost', 'carrier_total': 'wholesale_cost', 'net_charge': 'wholesale_cost',
    'residential_surcharge': 'residential_fee', 'residential_accessorial_fee': 'residential_fee'
}


class StatementLine(BaseModel):
    model_config = ConfigDict(extra='forbid', str_strip_whitespace=True)
    upstream_shipment_id: str = Field(min_length=1, max_length=200)
    currency: str
    invoice_weight_kg: Decimal = Field(ge=0, le=1000000, decimal_places=6)
    wholesale_cost: Decimal = Field(ge=0, le=10000000, decimal_places=2)
    residential_fee: Decimal = Field(ge=0, le=10000000, decimal_places=2)

    @field_validator('currency')
    @classmethod
    def currency_code(cls, value):
        value = value.upper()
        if value not in {'GBP', 'EUR', 'USD'}:
            raise ValueError('Unsupported currency')
        return value

    @field_validator('upstream_shipment_id')
    @classmethod
    def printable(cls, value):
        if any(ord(c) < 32 for c in value):
            raise ValueError('Invalid carrier reference')
        return value

    @field_validator('invoice_weight_kg', 'wholesale_cost', 'residential_fee', mode='before')
    @classmethod
    def exact_numbers(cls, value):
        return strict_decimal(str(value))


def parse_statement(raw, filename):
    suffix = PurePath(filename or '').suffix.lower()
    if suffix == '.csv':
        try:
            table = []
            for row in csv.reader(io.StringIO(raw.decode('utf-8-sig')), strict=True):
                table.append(row)
                if len(table) > MAX_ROWS + 1:
                    raise HTTPException(422, 'Statement exceeds 1000 rows')
        except (UnicodeError, csv.Error):
            raise HTTPException(422, 'Use a UTF-8 CSV with the documented column headers') from None
    elif suffix == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 200 or sum(f.file_size for f in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValueError('Expanded workbook too large')
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) != 1:
                    raise ValueError('Use one worksheet')
                sheet = workbook.active
                if sheet.max_column > 100 or sheet.max_row > MAX_ROWS + 1:
                    raise ValueError('Workbook dimensions exceed limit')
                table = []
                for row in sheet.iter_rows():
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('Formula cells are not invoice evidence')
                    table.append([cell.value for cell in row])
                    if len(table) > MAX_ROWS + 1:
                        raise ValueError('Too many rows')
            finally:
                workbook.close()
        except (ValueError, KeyError, TypeError, zipfile.BadZipFile, OSError, DefusedXmlException, ParseError):
            raise HTTPException(422, 'Use a single-sheet XLSX containing values, no formulas, and at most 1000 invoice rows') from None
    else:
        raise HTTPException(415, 'Upload a CSV or XLSX statement; legacy XLS files must be converted to XLSX')
    if len(table) < 2 or len(table) > MAX_ROWS + 1:
        raise HTTPException(422, 'Supply between 1 and 1000 statement rows')
    headers = [re.sub(r'[^a-z0-9]+', '_', str(h or '').strip().lower()).strip('_') for h in table[0]]
    headers = [HEADER_ALIASES.get(header, header) for header in headers]
    if len(headers) > 100 or len(set(headers)) != len(headers) or not HEADERS.issubset(headers):
        raise HTTPException(422, 'Required columns: ' + ', '.join(sorted(HEADERS)))
    result = []
    references = set()
    for index, cells in enumerate(table[1:], 2):
        if len(cells) != len(headers):
            raise HTTPException(422, f'Row {index} has an incorrect column count')
        try:
            line = StatementLine.model_validate({header: cell for header, cell in zip(headers, cells) if header in HEADERS})
        except ValueError:
            raise HTTPException(422, f'Row {index} has an invalid reference, currency or numeric value') from None
        if line.upstream_shipment_id in references:
            raise HTTPException(422, f'Row {index} repeats a carrier reference; consolidate split invoice lines first')
        references.add(line.upstream_shipment_id)
        result.append(line)
    return result


def compare_line(line, shipment):
    metadata = shipment['logistics_metadata']
    if isinstance(metadata, str):
        metadata = json.loads(metadata)
    if not isinstance(metadata,dict):
        metadata = {}
    verification = metadata.get('address_verification', {})
    if not isinstance(verification,dict):
        verification = {}
    commercial = verification.get('verified') is True and verification.get('is_commercial') is True
    findings = []
    packages=shipment.get('multi_parcel_dimensions') or []
    if isinstance(packages,str):packages=json.loads(packages)
    packed_weight=sum((Decimal(str(box['actual_weight_kg'])) for box in packages),Decimal(0)) if packages else shipment['parcel_weight_kg']
    weight_delta = abs(line.invoice_weight_kg - packed_weight)
    if weight_delta > Decimal('0.1'):
        findings.append({'code': 'WEIGHT_DELTA', 'difference_kg': str(weight_delta)})
    residential = line.residential_fee if commercial else Decimal('0')
    if residential > 0:
        findings.append({'code': 'RESIDENTIAL_FEE_ON_VERIFIED_COMMERCIAL_ADDRESS', 'amount': str(residential)})
    price_delta = max(Decimal('0'), line.wholesale_cost - shipment['wholesale_cost'])
    if price_delta > 0:
        findings.append({'code': 'WHOLESALE_COST_EXCEEDS_SNAPSHOT', 'amount': str(price_delta)})
    # Residential charges may already be inside the wholesale total: never add twice.
    estimated = max(price_delta, residential)
    baseline = {'parcel_weight_kg': str(packed_weight), 'wholesale_cost': str(shipment['wholesale_cost']),
                'currency': shipment['currency'], 'commercial_address_verified': commercial,
                'address_verification': verification}
    draft = 'Subject: Invoice review request ' + line.upstream_shipment_id + '\n\n'
    draft += 'Please review the following differences against our stored booking snapshot.\n'
    draft += '\n'.join(finding['code'] + ': ' + str(finding.get('amount', finding.get('difference_kg'))) for finding in findings)
    draft += '\nEstimated monetary discrepancy: ' + format(estimated, '.2f') + ' ' + line.currency
    draft += '\nPlease provide the carrier measurement and contractual charge evidence. This draft is not an adjudicated refund claim.'
    return findings, baseline, estimated, draft


def build_router(authenticate):
    router = APIRouter(prefix='/api/v1/finance', tags=['Statement audit'])

    @router.post('/audit-invoice-upload')
    async def upload(request: Request, file: UploadFile = File(...), invoice_reference: str = Form(...),
                     carrier_account_id: UUID | None = Form(default=None), context: Any = Depends(authenticate)):
        if context.role not in {'OWNER', 'ADMIN', 'ACCOUNTANT'}:
            raise HTTPException(403, 'Financial permission required')
        invoice_reference = invoice_reference.strip()
        if not invoice_reference or len(invoice_reference) > 100 or any(ord(c) < 32 for c in invoice_reference):
            raise HTTPException(422, 'Supply a printable invoice reference of at most 100 characters')
        try:
            raw = await file.read(MAX_BYTES + 1)
            if len(raw) > MAX_BYTES:
                raise HTTPException(413, 'Statement uploads are limited to 5 MB')
            lines = await asyncio.to_thread(parse_statement, raw, file.filename)
        finally:
            await file.close()
        results = []
        try:
            async with request.app.state.finance.pool.acquire() as connection:
                async with connection.transaction(isolation='repeatable_read'):
                    rows = await connection.fetch('''SELECT id,upstream_shipment_id,currency,wholesale_cost,
                        parcel_weight_kg,multi_parcel_dimensions,logistics_metadata,booked_at FROM public.shipments
                        WHERE tenant_id=$1 AND (upstream_shipment_id=ANY($2::text[]) OR logistics_metadata->'automation_label'->>'upstream_id'=ANY($2::text[]))
                        AND booked_at IS NOT NULL AND ($3::uuid IS NULL OR carrier_account_id=$3)''', context.tenant_id,
                        [line.upstream_shipment_id for line in lines], carrier_account_id)
                    for line in lines:
                        matches = []
                        for row in rows:
                            metadata = json.loads(row['logistics_metadata']) if isinstance(row['logistics_metadata'], str) else row['logistics_metadata']
                            if row['upstream_shipment_id'] == line.upstream_shipment_id or metadata.get('automation_label', {}).get('upstream_id') == line.upstream_shipment_id:
                                matches.append(row)
                        if len(matches) != 1 or not matches[0]['booked_at']:
                            results.append({'upstream_shipment_id': line.upstream_shipment_id, 'status': 'REVIEW_REQUIRED', 'reason': 'No unique confirmed booking matches this carrier reference'})
                            continue
                        shipment = matches[0]
                        if shipment['currency'] != line.currency:
                            results.append({'upstream_shipment_id': line.upstream_shipment_id, 'status': 'REVIEW_REQUIRED', 'reason': 'Invoice currency does not match booking currency'})
                            continue
                        engine = getattr(request.app.state, 'logistics_supervisor', None)
                        findings, baseline, estimated, draft = (engine.invoice.compare(line, shipment)
                            if engine is not None else compare_line(line, shipment))
                        data = line.model_dump(mode='json')
                        digest = hashlib.sha256(json.dumps(data, sort_keys=True).encode()).hexdigest()
                        dispute_id = None
                        if findings:
                            await connection.execute('''INSERT INTO public.billing_disputes(tenant_id,shipment_id,invoice_reference,
                                line_sha256,findings,baseline,invoice_line,confidence_level,estimated_overcharge,currency,email_draft,created_by)
                                VALUES($1,$2,$3,$4,$5::jsonb,$6::jsonb,$7::jsonb,100,$8,$9,$10,$11)
                                ON CONFLICT(tenant_id,invoice_reference,shipment_id) DO NOTHING''', context.tenant_id,
                                shipment['id'], invoice_reference, digest, json.dumps(findings), json.dumps(baseline),
                                json.dumps(data), estimated, line.currency, draft, context.user_id)
                            stored = await connection.fetchrow('SELECT id,line_sha256 FROM public.billing_disputes WHERE tenant_id=$1 AND invoice_reference=$2 AND shipment_id=$3', context.tenant_id, invoice_reference, shipment['id'])
                            if stored['line_sha256'] != digest:
                                raise HTTPException(409, 'This invoice was already audited with different line values; use a revised invoice reference')
                            dispute_id = str(stored['id'])
                        results.append({'upstream_shipment_id': line.upstream_shipment_id, 'shipment_id': str(shipment['id']),
                            'status': 'FLAGGED' if findings else 'MATCHED', 'confidence_level': 100,
                            'confidence_scope': 'Exact comparisons to stored evidence; not certainty of carrier liability',
                            'findings': findings, 'estimated_overcharge': format(estimated, '.2f'), 'currency': line.currency,
                            'dispute_id': dispute_id, 'email_draft': draft if findings else None})
        except asyncpg.PostgresError:
            raise HTTPException(503, 'Invoice audit could not be committed; retry the same statement') from None
        return {'invoice_reference': invoice_reference, 'results': results, 'emails_sent': False}

    @router.get('/billing-disputes')
    async def disputes(request: Request, context: Any = Depends(authenticate)):
        if context.role not in {'OWNER', 'ADMIN', 'ACCOUNTANT'}:
            raise HTTPException(403, 'Financial permission required')
        async with request.app.state.finance.pool.acquire() as connection:
            rows = await connection.fetch('SELECT id,shipment_id,invoice_reference,findings,estimated_overcharge,currency,email_draft,status,created_at FROM public.billing_disputes WHERE tenant_id=$1 ORDER BY created_at DESC LIMIT 100', context.tenant_id)
        return {'disputes': [dict(row) for row in rows]}

    return router
